"""Office regressions use mutated real packages and independent library output."""
from __future__ import annotations

import base64
import csv
from decimal import Decimal
from pathlib import Path
import re
import shutil
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from counso_bench.io import ROOT, BenchError
from counso_bench.office_check import (
    UnsupportedFormula, check_sales_workbook, sumifs, translate_formula,
)
from counso_bench.office_package import read_office_parts
from counso_bench.checks import workbook_cells


class OfficeV2Tests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def candidate(self, name='formulas.xlsx', mutate=None):
        path = self.directory / name
        shutil.copyfile(ROOT / 'gold/reference' / name, path)
        if mutate:
            with zipfile.ZipFile(path) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            mutate(parts)
            with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
                for name, body in parts.items():
                    archive.writestr(name, body)
        return path

    @staticmethod
    def change_formula(parts, callback):
        body = parts['xl/worksheets/sheet1.xml'].decode()
        body = re.sub(r'<f>(.*?)</f>', lambda match: '<f>' + callback(match[1]) + '</f>', body)
        parts['xl/worksheets/sheet1.xml'] = body.encode()

    def test_reference_packages_still_pass(self):
        for name, options in [
            ('chart.xlsx', {'chart_required': True}),
            ('formulas.xlsx', {'formula_required': True}),
            ('next-month.xlsx', {'with_april': True, 'formula_required': True, 'chart_required': True, 'dynamic_required': True}),
        ]:
            with self.subTest(name=name):
                path = self.candidate(name)
                self.assertTrue(read_office_parts(path))
                self.assertTrue(check_sales_workbook(path, ROOT, **options)[0])

    def test_undefined_structured_table_is_not_fixture_data(self):
        def mutate(parts):
            self.change_formula(parts, lambda value: value.replace('Sales!$C$2:$C$10', 'Sales[Amount]').replace('Sales!$A$2:$A$10', 'Sales[Month]').replace('Sales!$D$2:$D$10', 'Sales[Type]'))
        path = self.candidate(mutate=mutate)
        with self.assertRaisesRegex(BenchError, '未定义的原生表'):
            check_sales_workbook(path, ROOT, formula_required=True)

    def test_real_range_shape_mismatch_rejected(self):
        path = self.candidate(mutate=lambda parts: self.change_formula(parts, lambda value: value.replace('Sales!$A$2:$A$10', 'Sales!$A$2:$A$11')))
        with self.assertRaisesRegex(BenchError, '行列数不一致'):
            check_sales_workbook(path, ROOT, formula_required=True)

    def test_ranges_read_actual_candidate_cells(self):
        path = self.candidate()
        cells = workbook_cells(path)
        formula = cells['Summary!B2']['formula']
        self.assertEqual(sumifs(formula, cells)[0], Decimal('2000'))
        # The fixture argument is deliberately unchanged; candidate values decide.
        cells['Sales!C2']['value'] = str(Decimal(cells['Sales!C2']['value']) + 1)
        self.assertEqual(sumifs(formula, cells, rows=[{'Amount': '999999'}])[0], Decimal('2001'))

    def test_qualified_month_and_equality_criterion_are_accepted(self):
        def mutate(parts):
            self.change_formula(parts, lambda value: re.sub(r',A([2-4]),', r",'Summary'!$A\1,", value).replace('"SALE"', '"=SALE"'))
        path = self.candidate(mutate=mutate)
        self.assertTrue(check_sales_workbook(path, ROOT, formula_required=True)[0])

    def test_whole_column_shapes_compute_without_materializing_million_rows(self):
        path = self.candidate()
        cells = workbook_cells(path)
        formula = 'SUMIFS(Sales!C:C,Sales!A:A,Summary!A2,Sales!D:D,"sale")'
        self.assertEqual(sumifs(formula, cells)[0], Decimal('2000'))

    def test_known_formula_engine_limit_remains_unknown(self):
        path = self.candidate(mutate=lambda parts: self.change_formula(parts, lambda value: 'IFNA(' + value + ',0)'))
        with patch.dict('os.environ', {'COUNSO_SOFFICE': '/nonexistent/counso-office-engine'}):
            with self.assertRaises(UnsupportedFormula):
                check_sales_workbook(path, ROOT, formula_required=True)

    def test_next_month_absolute_row_lock_fails_real_copy(self):
        for pattern in ('$A$\\1', 'A$\\1'):
            with self.subTest(pattern=pattern):
                def mutate(parts):
                    self.change_formula(parts, lambda value: re.sub(r',A([2-5]),', lambda match: ',' + pattern.replace('\\1', match[1]) + ',', value))
                path = self.candidate('next-month.xlsx', mutate)
                passed, reason = check_sales_workbook(path, ROOT, with_april=True, formula_required=True, chart_required=True, dynamic_required=True)
                self.assertFalse(passed)
                self.assertIn('复制公式', reason)

    def test_next_month_relative_row_and_absolute_column_pass(self):
        path = self.candidate('next-month.xlsx', lambda parts: self.change_formula(parts, lambda value: re.sub(r',A([2-5]),', r',Summary!$A\1,', value)))
        self.assertTrue(check_sales_workbook(path, ROOT, with_april=True, formula_required=True, chart_required=True, dynamic_required=True)[0])

    def test_formula_translation_preserves_constants_tables_and_locks(self):
        formula = 'SUMIFS(Sales[Amount],Sales[Month],\'Q1\'!$A5,Sales[Type],"A5")+$B$2+B$3+$C4'
        expected = 'SUMIFS(Sales[Amount],Sales[Month],\'Q1\'!$A6,Sales[Type],"A5")+$B$2+B$3+$C5'
        self.assertEqual(translate_formula(formula, row_delta=1), expected)

    def test_missing_content_types_or_root_relationship_is_invalid(self):
        for missing in ('[Content_Types].xml', '_rels/.rels'):
            with self.subTest(missing=missing):
                path = self.candidate('chart.xlsx', lambda parts: parts.pop(missing))
                with self.assertRaises(BenchError):
                    read_office_parts(path)

    def test_unused_content_type_override_does_not_invent_a_part(self):
        def mutate(parts):
            parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace(b'</Types>', b'<Override PartName="/xl/worksheets/unused.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        path = self.candidate('chart.xlsx', mutate)
        parts = read_office_parts(path)
        self.assertNotIn('xl/worksheets/unused.xml', parts)
        self.assertTrue(check_sales_workbook(path, ROOT, chart_required=True)[0])
        # A declaration cannot excuse a missing actual relationship target.
        with zipfile.ZipFile(path) as archive:
            body = {name: archive.read(name) for name in archive.namelist()}
        body.pop('xl/worksheets/sheet1.xml')
        with zipfile.ZipFile(path, 'w') as archive:
            for name, data in body.items():
                archive.writestr(name, data)
        with self.assertRaises(BenchError):
            read_office_parts(path)

    def test_content_type_and_relationship_targets_are_checked(self):
        mutations = [
            lambda parts: parts.update({'[Content_Types].xml': parts['[Content_Types].xml'].replace(b'spreadsheetml.sheet.main+xml', b'spreadsheetml.invalid+xml')}),
            lambda parts: parts.update({'xl/_rels/workbook.xml.rels': parts['xl/_rels/workbook.xml.rels'].replace(b'Target="styles.xml"', b'Target="missing.xml"')}),
            lambda parts: parts.update({'xl/_rels/workbook.xml.rels': parts['xl/_rels/workbook.xml.rels'].replace(b'Target="styles.xml"', b'Target="../../outside.xml"')}),
        ]
        for mutate in mutations:
            with self.subTest(mutation=mutate):
                path = self.candidate('chart.xlsx', mutate)
                with self.assertRaises(BenchError):
                    read_office_parts(path)

    def test_media_relationship_targets_use_the_whole_package(self):
        def mutate(parts):
            parts['xl/media/image1.png'] = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScLbtAAAAABJRU5ErkJggg==')
            parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace(b'</Types>', b'<Default Extension="png" ContentType="image/png"/></Types>')
            parts['xl/drawings/_rels/drawing1.xml.rels'] = parts['xl/drawings/_rels/drawing1.xml.rels'].replace(b'</Relationships>', b'<Relationship Id="rIdImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>')
        path = self.candidate('chart.xlsx', mutate)
        self.assertIn('xl/media/image1.png', read_office_parts(path))
        # Business validation must also accept media relationships.
        self.assertTrue(check_sales_workbook(path, ROOT, chart_required=True)[0])

    def independent_openpyxl(self, with_image=False):
        try:
            import openpyxl
            from openpyxl.chart import BarChart, Reference
        except ImportError:
            self.skipTest('independent openpyxl control requires optional test dependency')
        book = openpyxl.Workbook(); summary = book.active; summary.title = 'Summary'
        summary.append(['Month', 'Revenue'])
        for record in [('2026-01', 2000), ('2026-02', 2400), ('2026-03', 3000)]:
            summary.append(record)
        sales = book.create_sheet('Sales'); sales.append(['Month', 'Region', 'Amount', 'Type'])
        with (ROOT / 'fixtures/sales.csv').open(newline='', encoding='utf-8') as source:
            for record in csv.DictReader(source):
                sales.append([record['Month'], record['Region'], float(record['Amount']), record['Type']])
        chart = BarChart(); chart.add_data(Reference(summary, min_col=2, min_row=1, max_row=4), titles_from_data=True)
        chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=4)); summary.add_chart(chart, 'D2')
        if with_image:
            try:
                from PIL import Image
                from openpyxl.drawing.image import Image as WorkbookImage
            except ImportError:
                self.skipTest('image control requires optional Pillow test dependency')
            image_path = self.directory / 'independent.png'; Image.new('RGB', (8, 8), 'blue').save(image_path)
            summary.add_image(WorkbookImage(str(image_path)), 'D20')
        path = self.directory / 'independent.xlsx'; book.save(path)
        loaded = openpyxl.load_workbook(path)
        self.assertEqual(len(loaded['Summary']._charts), 1)
        loaded.close()
        return path

    def test_independent_native_numref_chart_is_accepted(self):
        path = self.independent_openpyxl()
        parts = read_office_parts(path)
        chart_xml = next(body for name, body in parts.items() if name.startswith('xl/charts/') and name.endswith('.xml'))
        self.assertIn(b'<numRef>', chart_xml)
        self.assertTrue(check_sales_workbook(path, ROOT, chart_required=True)[0])

    def test_independent_native_chart_with_real_image_is_accepted(self):
        path = self.independent_openpyxl(with_image=True)
        self.assertTrue(any(name.startswith('xl/media/') for name in read_office_parts(path)))
        self.assertTrue(check_sales_workbook(path, ROOT, chart_required=True)[0])


if __name__ == '__main__':
    unittest.main()
